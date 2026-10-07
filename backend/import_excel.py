"""One-time import of JATIM, TG and Job Open sheets. Usage: python import_excel.py [--force]"""
import asyncio
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from dotenv import load_dotenv

load_dotenv(Path(__file__).parent / ".env")

import openpyxl  # noqa: E402
from core import db  # noqa: E402

FILE = Path(__file__).parent / "data" / "jurnal_hakusa.xlsx"
BASE_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)
_seq = 0


def stamp() -> str:
    global _seq
    _seq += 1
    return (BASE_TIME + timedelta(microseconds=_seq)).isoformat()


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) and v != 0 else 0.0


def dt(v):
    return v.date().isoformat() if isinstance(v, datetime) else None


def has(text, words):
    t = text.lower()
    return any(re.search(w, t) for w in words)


def cat_keluar(text):
    if has(text, [r"visa"]):
        return "Visa"
    if has(text, [r"ktkln"]):
        return "KTKLN"
    if has(text, [r"mcu"]):
        return "MCU"
    if has(text, [r"tsk", r"spj", r"\bpmi\b", r"\blpk\b", r"bayar ke", r"\btrf\b", r"endo", r"hibiki", r"\bhim\b",
                  r"enlink", r"inji", r"human"]):
        return "TSK"
    if has(text, [r"tiket", r"\bbis\b", r"\bbus\b", r"grab", r"pesawat", r"transport", r"hotel", r"ongkir"]):
        return "Transport & Akomodasi"
    if has(text, [r"gaji", r"bayaran"]):
        return "Gaji"
    return None


STATUS_WORDS = ("lunas", "selesai", "selsai", "lunas selsai")


def _cell(r, i):
    return r[i] if i < len(r) else None


def _is_header(r, c):
    a, m, k = _cell(r, c), _cell(r, c + 2), _cell(r, c + 3)
    return (isinstance(a, str) and a.strip().lower() == "tanggal") or \
        (isinstance(m, str) and m.strip().lower() == "masuk" and isinstance(k, str) and k.strip().lower() == "keluar")


def _is_name(r, c):
    v = _cell(r, c + 1)
    return isinstance(v, str) and v.strip() and v.strip().lower() not in STATUS_WORDS + ("jenis transaksi",) \
        and not num(_cell(r, c + 2)) and not num(_cell(r, c + 3)) and _cell(r, c + 4) is None


def _segments(rows):
    groups = sorted({c for r in rows for c in range(len(r)) if _is_header(r, c)})
    blocks = []
    for c in groups:
        cur = None
        for i, r in enumerate(rows):
            if _is_header(r, c):
                if cur and cur["header"] is None and not cur["rows"] and i - cur["start"] <= 3:
                    cur["header"] = i
                else:
                    cur = {"name": f"Tanpa Nama (baris {i + 1})", "start": i, "header": i, "rows": [], "c": c}
                    blocks.append(cur)
            elif _is_name(r, c):
                cur = {"name": r[c + 1].strip(), "start": i, "header": None, "rows": [], "c": c}
                blocks.append(cur)
            elif cur:
                cur["rows"].append(r)
    return blocks


def parse_tg(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    result = {}
    for b in _segments(rows):
        c, name = b["c"], b["name"]
        entries, lunas, last_date = [], False, None
        block_dates = [dt(_cell(r, c)) for r in b["rows"] if dt(_cell(r, c))]
        for r in b["rows"]:
            r = r + [None] * 6
            desc = r[c + 1].strip() if isinstance(r[c + 1], str) else ""
            masuk, keluar, saldo = num(r[c + 2]), num(r[c + 3]), r[c + 4]
            if desc.lower() in STATUS_WORDS:
                lunas = True
            d = dt(r[c])
            if d:
                last_date = d
            if (masuk <= 0 and keluar <= 0) or (masuk < 0 or keluar < 0):
                continue
            if not desc and saldo is None:
                continue
            entries.append({"tanggal": d or last_date or (block_dates[0] if block_dates else ""),
                            "keterangan": desc or "(tanpa keterangan)", "masuk": masuk, "keluar": keluar})
        if not entries and not b["header"]:
            continue
        key = re.sub(r"\s+", " ", name.lower())
        cur = result.setdefault(key, {"nama": re.sub(r"\s+", " ", name), "entries": [], "lunas": False})
        cur["entries"].extend(entries)
        cur["lunas"] = cur["lunas"] or lunas
    return list(result.values())


def parse_jatim(ws, student_firsts):
    out, last_date = [], ""
    for r in list(ws.iter_rows(values_only=True))[2:]:
        r = list(r) + [None] * 8
        d = dt(r[0])
        if d:
            last_date = d
        keluar, masuk = num(r[4]), num(r[5])
        if masuk <= 0 and keluar <= 0:
            continue
        desc = str(r[2]).strip() if r[2] is not None else ""
        desc = desc or "(tanpa keterangan)"
        if masuk > 0:
            if has(desc, [r"\bles\b", r"kelas", r"\bn4\b", r"\bn5\b", r"smk", r"kursus"]):
                kat = "Kursus/Kelas N4-N5"
            elif any(f in desc.lower() for f in student_firsts):
                kat = "Job Matching"
            else:
                kat = "Lainnya"
        else:
            kat = cat_keluar(desc) or "Operasional"
        out.append({"tanggal": d or last_date, "keterangan": desc, "ref": str(r[3]) if isinstance(r[3], str) else "",
                    "masuk": masuk if masuk > 0 else 0.0, "keluar": keluar if keluar > 0 else 0.0, "kategori": kat})
    return out


def parse_jobs(ws):
    jobs = []
    for r in list(ws.iter_rows(values_only=True))[3:]:
        r = list(r) + [None] * 9
        if not isinstance(r[0], str) or not r[0].strip() or r[0].lower().startswith("target"):
            continue
        jobs.append({"nama_job": r[0].strip(), "tsk": (r[1] or "").strip() if isinstance(r[1], str) else "",
                     "status": (r[2] or "").strip() if isinstance(r[2], str) else "",
                     "jumlah_peserta": int(num(r[3])), "lolos": int(num(r[4])), "tidak_lolos": int(num(r[5])),
                     "biaya": num(r[6]), "catatan": ""})
    return jobs


async def main(force: bool):
    if await db.transactions.count_documents({}) and not force:
        print("Data sudah ada. Pakai --force untuk menimpa.")
        return
    if force:
        await db.transactions.delete_many({})
        await db.students.delete_many({})
        await db.jobs.delete_many({})
    wb = openpyxl.load_workbook(FILE, data_only=True)
    students = parse_tg(wb["TG"])
    tg_count = 0
    for s in students:
        res = await db.students.insert_one({
            "nama": s["nama"], "job": "", "tsk": "", "no_hp": "", "tanggal_daftar": "", "biaya": 0,
            "est_visa": 2300000, "est_ktkln": 200000, "est_mcu": 800000, "est_tsk": 15000000,
            "status": "Selesai" if s["lunas"] else "Proses",
            "catatan": "Diimpor dari sheet TG. Isi biaya job matching untuk menghitung kekurangan.",
            "created_at": stamp()})
        sid = str(res.inserted_id)
        for e in s["entries"]:
            kat = "Job Matching" if e["masuk"] > 0 else (cat_keluar(e["keterangan"]) or "Lainnya")
            await db.transactions.insert_one({**e, "ref": "", "kategori": kat, "buku": "Kas TG", "student_id": sid,
                                              "source": "import_tg", "created_at": stamp()})
            tg_count += 1
    firsts = {s["nama"].lower().split()[0] for s in students if len(s["nama"].split()[0]) >= 4}
    jatim = parse_jatim(wb["JATIM"], firsts)
    for e in jatim:
        await db.transactions.insert_one({**e, "buku": "Kas JATIM", "student_id": None, "source": "import_jatim",
                                          "created_at": stamp()})
    jobs = parse_jobs(wb["Sheet7"])
    if jobs:
        await db.jobs.insert_many(jobs)
    print(f"Siswa TG: {len(students)}, transaksi TG: {tg_count}, transaksi JATIM: {len(jatim)}, job: {len(jobs)}")


if __name__ == "__main__":
    asyncio.run(main("--force" in sys.argv))
