import io
from collections import defaultdict
from datetime import date

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from core import get_current_user
from routes_finance import ledger
from routes_students import all_students

router = APIRouter(prefix="/api", dependencies=[Depends(get_current_user)])


@router.get("/dashboard")
async def dashboard(buku: str = ""):
    rows = await ledger({"buku": buku} if buku else {})
    monthly = defaultdict(lambda: {"masuk": 0.0, "keluar": 0.0})
    inc_cat, exp_cat = defaultdict(float), defaultdict(float)
    for r in rows:
        m = r["tanggal"][:7]
        if m:
            monthly[m]["masuk"] += r["masuk"]
            monthly[m]["keluar"] += r["keluar"]
        if r["masuk"]:
            inc_cat[r["kategori"] or "Tanpa Kategori"] += r["masuk"]
        if r["keluar"]:
            exp_cat[r["kategori"] or "Tanpa Kategori"] += r["keluar"]
    this_month = date.today().isoformat()[:7]
    months = sorted(m for m in monthly if m <= this_month)[-12:]
    students = await all_students()
    with_fee = [s for s in students if s["biaya"] > 0]
    pending = sorted([s for s in with_fee if s["kekurangan"]], key=lambda s: -s["kekurangan"])
    return {
        "total_masuk": sum(r["masuk"] for r in rows),
        "total_keluar": sum(r["keluar"] for r in rows),
        "saldo": rows[-1]["saldo"] if rows else 0,
        "jumlah_transaksi": len(rows),
        "monthly": [{"bulan": m, **monthly[m]} for m in months],
        "income_by_category": [{"kategori": k, "total": v} for k, v in sorted(inc_cat.items(), key=lambda x: -x[1])],
        "expense_by_category": [{"kategori": k, "total": v} for k, v in sorted(exp_cat.items(), key=lambda x: -x[1])],
        "recent": rows[::-1][:8],
        "siswa": {
            "jumlah": len(students),
            "lunas": sum(1 for s in students if s["lunas"]),
            "belum_lunas": len(pending),
            "biaya_belum_diatur": len(students) - len(with_fee),
            "total_tagihan": sum(s["biaya"] for s in with_fee),
            "total_bayar": sum(s["total_bayar"] for s in students),
            "total_keluar": sum(s["total_keluar"] for s in students),
            "total_kekurangan": sum(s["kekurangan"] or 0 for s in with_fee),
            "profit_aktual": sum(s["profit_aktual"] for s in students),
            "top_kekurangan": [{"id": s["id"], "nama": s["nama"], "biaya": s["biaya"],
                                "total_bayar": s["total_bayar"], "kekurangan": s["kekurangan"]} for s in pending[:6]],
        },
    }


HEAD_FILL = PatternFill("solid", fgColor="1E3A8A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0;[Red]-#,##0'


def _sheet(wb, title, headers, data, money_cols=(), widths=None):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center")
    for row in data:
        ws.append(row)
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = MONEY
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    return ws


@router.get("/export/xlsx")
async def export_xlsx(buku: str = ""):
    rows = await ledger({"buku": buku} if buku else {})
    students = await all_students()
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Jurnal Umum", ["Tanggal", "Ref", "Item / Keterangan", "Kategori", "Buku Kas", "Siswa",
                               "Jurnal Masuk", "Jurnal Keluar", "Saldo"],
           [[r["tanggal"], r["ref"], r["keterangan"], r["kategori"], r["buku"], r["student_nama"] or "",
             r["masuk"] or None, r["keluar"] or None, r["saldo"]] for r in rows],
           money_cols=(7, 8, 9), widths=[12, 10, 45, 20, 14, 22, 16, 16, 16])
    by_cat = defaultdict(list)
    for r in rows:
        if r["masuk"] > 0:
            by_cat[r["kategori"] or "Tanpa Kategori"].append(r)
    _sheet(wb, "Rekap Pemasukan", ["Jenis Pemasukan", "Jumlah Transaksi", "Total"],
           [[k, len(v), sum(x["masuk"] for x in v)] for k, v in by_cat.items()], money_cols=(3,), widths=[30, 18, 18])
    for k, v in by_cat.items():
        name = "Masuk - " + "".join(ch for ch in k if ch not in "[]:*?/\\")
        _sheet(wb, name, ["Tanggal", "Keterangan", "Siswa", "Buku Kas", "Jumlah"],
               [[x["tanggal"], x["keterangan"], x["student_nama"] or "", x["buku"], x["masuk"]] for x in v],
               money_cols=(5,), widths=[12, 45, 22, 14, 16])
    _sheet(wb, "Siswa TG", ["Nama", "Job", "TSK", "Status", "Biaya Job Matching", "Total Bayar", "Kekurangan",
                            "Total Pengeluaran", "Profit Aktual", "Estimasi Profit", "Lunas"],
           [[s["nama"], s["job"], s["tsk"], s["status"], s["biaya"] or None, s["total_bayar"], s["kekurangan"],
             s["total_keluar"], s["profit_aktual"], s["estimasi_profit"], "Ya" if s["lunas"] else "Belum"]
            for s in students], money_cols=(5, 6, 7, 8, 9, 10), widths=[28, 22, 18, 12, 18, 16, 16, 18, 16, 16, 8])
    names = {s["id"]: s["nama"] for s in students}
    student_rows = sorted([r for r in rows if r["student_id"]], key=lambda r: (names.get(r["student_id"], ""), r["tanggal"]))
    _sheet(wb, "Rincian Siswa TG", ["Siswa", "Tanggal", "Keterangan", "Kategori", "Masuk", "Keluar"],
           [[names.get(r["student_id"], ""), r["tanggal"], r["keterangan"], r["kategori"], r["masuk"] or None,
             r["keluar"] or None] for r in student_rows], money_cols=(5, 6), widths=[28, 12, 40, 20, 16, 16])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Jurnal_Hakusa_{date.today().isoformat()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})
